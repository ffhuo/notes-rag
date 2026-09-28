"""解析器·Excel — .xlsx / .xlsm 转 Markdown（docs/design.md §16.1）。

能力：
- openpyxl 逐 sheet 读取，输出「`# sheet名` 标题 + GFM pipe table」
- 首行作表头，其后补分隔行 —— 与 docx 的表格输出规范一致，chunker 的表格保护
  （`_MD_TABLE_RE`）才认得出；超长表格切分时每个子表都会重复表头
- 单元格规整：None → 空串、`|` → `\\|`、换行压成空格；每行补齐到最宽行的列数
  （ragged row 会让 GFM 表格错位）

取舍：
- data_only=True：取公式的**计算结果**而非 `=SUM(...)` 公式串（索引公式文本没有检索价值；
  若工作簿从未被 Excel 计算并保存过，公式单元格会读到 None）
- read_only=True：流式读取，避免大表把整簿载入内存
- 不支持老式 .xls：openpyxl 只能读 OOXML（.xlsx/.xlsm），故 supported_exts 不含 .xls

关联方案：docs/design.md §16（多格式文件与过滤设计）。
"""
from __future__ import annotations

import os
from collections.abc import Iterator
from pathlib import Path

import openpyxl

from app.parsers.base import DocumentParser, ParsedDocument
from app.parsers.registry import register


def _cell_text(value: object) -> str:
    """单元格值 → 单行文本：None 归一为空串，换行压成空格，`|` 转义。

    `|` 不转义会多切出一列、整张表错位；单元格内换行不压平会让表格断行。
    """
    if value is None:
        return ""
    text = str(value).strip()
    for ch in ("\r\n", "\n", "\r"):
        text = text.replace(ch, " ")
    return text.replace("|", "\\|")


def _iter_text_rows(sheet: openpyxl.worksheet.worksheet.Worksheet) -> Iterator[list[str]]:
    """逐行产出单元格文本，跳过全空行。

    全空行必须丢弃：GFM 表格里夹一个空行会把表格截断，后半段退化成普通文本。
    """
    for row in sheet.iter_rows(values_only=True):
        cells = [_cell_text(value) for value in row]
        if any(cells):
            yield cells


def _render_table(rows: list[list[str]]) -> str:
    """二维单元格 → GFM pipe table：首行表头 + 分隔行 + 数据行。

    每行补齐到最宽行的列数：数据行列数与表头不一致时，GFM 渲染会错位或丢列。
    """
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    lines = [
        "| " + " | ".join(padded[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in padded[1:])
    return "\n".join(lines)


class ExcelParser(DocumentParser):
    supported_exts = (".xlsx", ".xlsm")

    def parse(self, path: Path) -> ParsedDocument:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            blocks: list[str] = []
            for sheet in workbook.worksheets:
                rows = list(_iter_text_rows(sheet))
                if not rows:
                    continue                    # 空 sheet 不产出标题，免得留下只有标题的空分块
                blocks.append(f"# {sheet.title}")
                blocks.append(_render_table(rows))
            content = "\n\n".join(blocks)
            meta = {"source_type": "excel", "sheets": list(workbook.sheetnames)}
        finally:
            # read_only 模式必须显式关闭，否则底层 zip 句柄要到 GC 才释放
            workbook.close()

        return ParsedDocument(
            content=content,
            title=path.stem,
            mtime=os.path.getmtime(path),
            meta=meta,
        )


register(ExcelParser())
