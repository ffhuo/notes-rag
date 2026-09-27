"""Word(.docx) 解析器 — 正文结构、内嵌图片引用与 data: 分支校验。

关注点：
- 正文按「文档顺序」输出，标题转 Markdown `#` 前缀（切分因此能按标题切段）
- 内嵌图片产出 ImageRef：正文留短标记作为定位锚点，真实字节走 Data URL（不落盘）
- 图片的质量门槛与本地图一致 —— 内嵌图若不过尺寸校验就保留原语法，不占用多模态调用

样本全部在 tmp_path 里**动态生成**（python-docx + Pillow），不往仓库塞二进制。

关联方案：docs/design.md §16；图片方案见 image-progress 约定。
"""
import base64
import hashlib
import io
from dataclasses import replace
from pathlib import Path

import docx
import pytest
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches
from PIL import Image

from app.parsers.base import ImageRef
from app.parsers.docx import DocxParser
from app.services import image_service

Marker = str


def _png_bytes(width: int = 200, height: int = 200, color=(10, 120, 200)) -> bytes:
    """造一张纯色 PNG（默认 200×200，刚好达到 image_min_width / height 门槛）。"""
    buf = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buf, format="PNG")
    return buf.getvalue()


def _make_docx(tmp_path: Path, with_image: bool = True, image_px: int = 200) -> Path:
    """生成一份最小 docx：一级标题 + 正文段 +（可选）内嵌图 + 2×2 表格。"""
    document = docx.Document()
    document.add_heading("架构笔记", level=1)
    document.add_paragraph("正文第一段。")
    if with_image:
        document.add_picture(io.BytesIO(_png_bytes(image_px, image_px)), width=Inches(2))
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "D"
    path = tmp_path / "sample.docx"
    document.save(str(path))
    return path


def test_parse_text_and_heading(tmp_path):
    parsed = DocxParser().parse(_make_docx(tmp_path))

    assert "# 架构笔记" in parsed.content          # 标题样式 → markdown 标题
    assert "正文第一段。" in parsed.content
    assert parsed.title == "sample"
    assert parsed.meta["source_type"] == "docx"


def test_registered_for_docx_extension():
    """注册必须生效：否则扫描到 .docx 也会被静默跳过（缺依赖时正是这种表现）。"""
    from app.parsers import registry

    assert isinstance(registry.get_parser(".docx"), DocxParser)


def test_parse_table_as_gfm_pipe_table(tmp_path):
    """表格必须是合法 GFM pipe table：首尾带 `|` 且表头后有分隔行。

    少了任一点，前端 marked 都不会渲染成 <table>（检索结果里只会看到一行竖线文本）。
    """
    parsed = DocxParser().parse(_make_docx(tmp_path, with_image=False))

    assert "| A | B |" in parsed.content
    assert "| --- | --- |" in parsed.content
    assert "| C | D |" in parsed.content


def test_table_pipe_in_cell_escaped(tmp_path):
    """单元格内的 `|` 必须转义，否则会多切出一列、整张表错位。"""
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A|B"
    table.cell(0, 1).text = "C"
    table.cell(1, 0).text = "D"
    table.cell(1, 1).text = "E"
    path = tmp_path / "pipe.docx"
    document.save(str(path))

    parsed = DocxParser().parse(path)

    assert r"| A\|B | C |" in parsed.content


def test_large_table_split_repeats_header(tmp_path):
    """大表格切分时重复表头 —— 只有被识别为 GFM 表格，chunker 的表格保护才生效。"""
    document = docx.Document()
    document.add_heading("数据表", level=1)
    table = document.add_table(rows=31, cols=2)
    for i, row in enumerate(table.rows):
        row.cells[0].text = f"字段{i}"
        row.cells[1].text = "值" * 40
    path = tmp_path / "big.docx"
    document.save(str(path))

    parsed = DocxParser().parse(path)
    chunks = parsed.to_chunks(fmt="markdown", max_chars=300)
    table_chunks = [chunk for chunk in chunks if "| --- |" in chunk.text]

    assert len(table_chunks) > 1, "表格未被识别为表格，切分保护失效"
    assert all("字段0" in chunk.text for chunk in table_chunks)  # 每个子表都带表头


CODE_LINES = ["def hello():", "    print('hi')", "    return 1"]
FENCED = "```\ndef hello():\n    print('hi')\n    return 1\n```"


def _add_mono(document, text: str):
    """加一个等宽字体段落（Word 里手打代码最常见的形态）。"""
    run = document.add_paragraph().add_run(text)
    run.font.name = "Consolas"
    return run


def test_mono_paragraphs_merged_into_fence(tmp_path):
    """连续的等宽段落必须合并成一个围栏块，且行间不插空行（否则缩进代码被拆散）。"""
    document = docx.Document()
    document.add_heading("代码示例", level=1)
    for line in CODE_LINES:
        _add_mono(document, line)
    path = tmp_path / "mono.docx"
    document.save(str(path))

    content = DocxParser().parse(path).content

    assert FENCED in content


def test_single_mono_paragraph_with_line_breaks_fenced(tmp_path):
    """一行 Word 段落里用软换行写多行代码（等宽字体）—— 同样要成块。"""
    document = docx.Document()
    run = document.add_paragraph().add_run("\n".join(CODE_LINES))
    run.font.name = "Consolas"
    path = tmp_path / "softbreak.docx"
    document.save(str(path))

    assert FENCED in DocxParser().parse(path).content


def test_code_style_paragraphs_fenced(tmp_path):
    """样式名命中预格式样式（此处为自建「Code」）时成块，与字体无关。"""
    document = docx.Document()
    document.add_heading("代码示例", level=1)
    style = document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    for line in CODE_LINES:
        document.add_paragraph(line, style=style)
    path = tmp_path / "styled.docx"
    document.save(str(path))

    assert FENCED in DocxParser().parse(path).content


def test_single_mono_line_stays_paragraph(tmp_path):
    """单个单行等宽段落是「正文里提到的一小段」的概率更高，不升级为代码块。"""
    document = docx.Document()
    document.add_paragraph("配置文件叫 config.yaml。")
    _add_mono(document, "print('hi')")
    path = tmp_path / "oneliner.docx"
    document.save(str(path))

    content = DocxParser().parse(path).content

    assert "```" not in content
    assert "print('hi')" in content


def test_code_block_keeps_blank_line_and_closes_before_heading(tmp_path):
    """代码段内的空行要保留，代码段结束于后续标题（标题不能被吞进围栏）。"""
    document = docx.Document()
    document.add_heading("代码示例", level=1)
    _add_mono(document, "def a():")
    _add_mono(document, "    return 1")
    document.add_paragraph()
    _add_mono(document, "def b():")
    _add_mono(document, "    return 2")
    document.add_heading("说明", level=1)
    path = tmp_path / "blank.docx"
    document.save(str(path))

    content = DocxParser().parse(path).content

    assert "```\ndef a():\n    return 1\n\ndef b():\n    return 2\n```" in content
    assert content.rstrip().endswith("# 说明")


def test_mono_heading_stays_heading(tmp_path):
    """等宽字体的标题仍是标题 —— 判据里必须先排除标题样式。"""
    document = docx.Document()
    paragraph = document.add_heading("代码示例", level=1)
    for run in paragraph.runs:
        run.font.name = "Consolas"
    path = tmp_path / "monoheading.docx"
    document.save(str(path))

    content = DocxParser().parse(path).content

    assert content.startswith("# 代码示例")
    assert "```" not in content


def test_code_containing_fence_uses_tilde(tmp_path):
    """代码正文自带 ``` 时改用 ~~~ 围栏，否则围栏会被提前闭合。"""
    document = docx.Document()
    _add_mono(document, "```python")
    _add_mono(document, "print(1)")
    path = tmp_path / "fence.docx"
    document.save(str(path))

    content = DocxParser().parse(path).content

    assert content.startswith("~~~\n")
    assert content.endswith("\n~~~")


def test_parse_embedded_image(tmp_path):
    parsed = DocxParser().parse(_make_docx(tmp_path))
    assert len(parsed.images) == 1
    ref = parsed.images[0]
    # 正文里留的是短标记（不是整段 base64），锚点必须与偏移自洽
    assert ref.raw in parsed.content
    assert parsed.content[ref.start : ref.end] == ref.raw
    # 真实字节走 Data URL，交给 image_service 的 data: 分支
    assert ref.target.startswith("data:image/png;base64,")
    decoded = base64.b64decode(ref.target.split(",", 1)[1])
    assert Image.open(io.BytesIO(decoded)).size == (200, 200)


def test_parse_without_image(tmp_path):
    parsed = DocxParser().parse(_make_docx(tmp_path, with_image=False))

    assert parsed.images == []
    assert "docx-media" not in parsed.content


def test_image_marker_survives_chunking(tmp_path):
    """切分后标记仍完整落在某个 chunk 里 —— attach_images 靠 ref.raw 定位，断了就回填不了。"""
    parsed = DocxParser().parse(_make_docx(tmp_path))
    chunks = parsed.to_chunks(fmt="markdown", max_chars=2000)

    assert any(ref.raw in chunk.text for chunk in chunks for ref in parsed.images)


@pytest.mark.asyncio
async def test_embedded_image_uses_local_image_standard():
    """内嵌图与本地图共用校验标准，并按原始字节做内容寻址。"""
    png = _png_bytes()
    ref = ImageRef(
        raw="![图片](docx-media:1)", start=0, end=1, alt="图片",
        target="data:image/png;base64," + base64.b64encode(png).decode(), kind="md",
    )

    prepared = await image_service.prepare(ref, Path("."), Path("."))

    assert prepared.ok
    assert prepared.source_kind == "embedded"
    assert prepared.source_ref == ""            # 内嵌图没有可回填的外部地址
    assert prepared.content_key == hashlib.sha256(png).hexdigest()   # 与本地图同规范
    assert (prepared.width, prepared.height, prepared.size_bytes) == (200, 200, len(png))
    assert prepared.payload_url.startswith("data:image/png;base64,")


@pytest.mark.asyncio
async def test_embedded_small_image_rejected():
    """小图标（< 200px）不合格：保留原语法，不送去多模态模型。"""
    small = _png_bytes(100, 100)
    ref = ImageRef(
        raw="![图片](docx-media:1)", start=0, end=1, alt="图片",
        target="data:image/png;base64," + base64.b64encode(small).decode(), kind="md",
    )

    prepared = await image_service.prepare(ref, Path("."), Path("."))

    assert not prepared.ok
    assert prepared.reason == "invalid_data_url"


@pytest.mark.asyncio
async def test_non_base64_data_url_rejected():
    """percent-encoded 的 data URL 不处理（拿不到可信字节，宁可不做也不送错的载荷）。"""
    ref = ImageRef(
        raw="![图片](docx-media:1)", start=0, end=1, alt="图片",
        target="data:image/svg+xml,<svg xmlns='http://www.w3.org/2000/svg'/>", kind="md",
    )

    prepared = await image_service.prepare(ref, Path("."), Path("."))

    assert not prepared.ok
