"""Excel(.xlsx) 解析器 — 多 sheet、GFM 表格输出与超长表格拆分。

关注点：
- 每个 sheet 转成「`# sheet名` + GFM pipe table」，标题让切分能按 sheet 分段并生成面包屑
- 表格必须是合法 GFM（首行表头 + 分隔行）：否则前端 marked 不渲染成 <table>，
  chunker 的表格保护（`_MD_TABLE_RE`）也认不出这张表
- 超长表格被切分时每个子表都要重复表头 —— 否则后半段子表失去列语义

样本全部在 tmp_path 里动态生成（openpyxl 写入），不往仓库塞二进制。

关联方案：docs/design.md §16。
"""
from pathlib import Path

import openpyxl

from app.parsers.excel import ExcelParser
from app.parsers.registry import get_parser


def _save(workbook: openpyxl.Workbook, tmp_path: Path, name: str = "sample.xlsx") -> Path:
    path = tmp_path / name
    workbook.save(str(path))
    return path


def _make_xlsx(tmp_path: Path) -> Path:
    """两个 sheet：预算表（表头 + 2 行数据）与备注（表头 + 1 行数据）。"""
    workbook = openpyxl.Workbook()
    budget = workbook.active
    budget.title = "预算表"
    budget.append(["项目", "金额"])
    budget.append(["服务器", 1200])
    budget.append(["域名", 80])
    notes = workbook.create_sheet("备注")
    notes.append(["说明"])
    notes.append(["仅记录"])
    return _save(workbook, tmp_path)


def test_registered_for_xlsx_extension():
    """注册必须生效：否则扫描到 .xlsx 会被静默跳过（缺依赖时正是这种表现）。"""
    assert isinstance(get_parser(".xlsx"), ExcelParser)


def test_old_xls_not_registered():
    """.xls 不在支持范围（openpyxl 读不了 OOXML 之前的格式），不能虚假声明。"""
    assert get_parser(".xls") is None


def test_parse_sheets_as_headings(tmp_path):
    """sheet 名转 Markdown 标题：切分才能按 sheet 分段，并让检索结果带上面包屑。"""
    parsed = ExcelParser().parse(_make_xlsx(tmp_path))

    assert "# 预算表" in parsed.content
    assert "# 备注" in parsed.content
    assert parsed.title == "sample"
    assert parsed.meta["source_type"] == "excel"
    assert parsed.meta["sheets"] == ["预算表", "备注"]


def test_table_is_gfm_pipe_table(tmp_path):
    """表格必须是合法 GFM：首尾带 `|` 且表头后紧跟分隔行。"""
    content = ExcelParser().parse(_make_xlsx(tmp_path)).content

    assert "| 项目 | 金额 |" in content
    assert "| --- | --- |" in content
    assert "| 服务器 | 1200 |" in content


def test_cell_pipe_escaped_and_newline_flattened(tmp_path):
    """单元格内的 `|` 要转义、换行要压平，否则多切一列 / 表格断行。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "特殊字符"
    sheet.append(["列A", "列B"])
    sheet.append(["a|b", "第一行\n第二行"])

    content = ExcelParser().parse(_save(workbook, tmp_path)).content

    assert r"| a\|b | 第一行 第二行 |" in content


def test_ragged_rows_padded_to_widest(tmp_path):
    """短行补齐到最宽行 —— 列数不齐会让 GFM 表格错位。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "不齐"
    sheet.append(["A", "B", "C"])
    sheet.append(["1"])

    content = ExcelParser().parse(_save(workbook, tmp_path)).content

    assert "| 1 |  |  |" in content


def test_none_cell_becomes_empty(tmp_path):
    """空单元格渲染为空串，不能让 "None" 落进索引。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "空值"
    sheet.append(["A", "B"])
    sheet.append(["1", None])

    content = ExcelParser().parse(_save(workbook, tmp_path)).content

    assert "| 1 |  |" in content
    assert "None" not in content


def test_blank_rows_skipped_and_empty_sheet_omitted(tmp_path):
    """全空行必须丢弃（夹在表格里会截断 GFM 表格）；空 sheet 不产出标题。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "有空行"
    sheet.append(["A"])
    sheet.append([None])
    sheet.append(["B"])
    workbook.create_sheet("空表")

    content = ExcelParser().parse(_save(workbook, tmp_path)).content

    assert "# 空表" not in content
    table_lines = [line for line in content.splitlines() if line.startswith("|")]
    assert table_lines == ["| A |", "| --- |", "| B |"]


def test_large_table_split_repeats_header(tmp_path):
    """超长表格切分时每个子表都带表头与分隔行 —— 这正是 xlsx 走 markdown 策略的目的。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "大表"
    sheet.append(["字段", "值"])
    for i in range(30):
        sheet.append([f"字段{i}", "值" * 40])
    path = _save(workbook, tmp_path, "big.xlsx")

    parsed = ExcelParser().parse(path)
    chunks = parsed.to_chunks(fmt="markdown", max_chars=300)
    table_chunks = [chunk for chunk in chunks if "| --- |" in chunk.text]

    assert len(table_chunks) > 1, "表格未被识别为表格，切分保护失效"
    for chunk in table_chunks:
        assert "| 字段 | 值 |" in chunk.text           # 表头随每块重复
        # 表头在前、分隔行紧随其后，顺序错了仍不是合法 GFM
        assert chunk.text.index("| --- |") > chunk.text.index("| 字段 | 值 |")
